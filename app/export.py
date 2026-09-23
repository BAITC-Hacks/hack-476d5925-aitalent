"""Экспорт протокола совещания в DOCX и PDF."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from .config import EXPORT_DIR, FONTS_DIR

CONSENT_NOTE = "Участники совещания уведомлены о ведении аудиозаписи и её обработке с помощью ИИ."


def _fmt_date(iso: str | None) -> str:
    if not iso:
        return "не указан"
    try:
        return date.fromisoformat(iso).strftime("%d.%m.%Y")
    except ValueError:
        return iso


def _ts(sec: float) -> str:
    m, s = divmod(int(sec), 60)
    return f"{m:02d}:{s:02d}"


def _speaker(label: str, names: dict) -> str:
    name = names.get(label)
    return f"{name}" if name else label


def _prepare(meeting: dict) -> dict:
    names = meeting.get("speakers") or {}
    speakers = []
    for s in meeting.get("segments", []):
        if s["speaker"] not in speakers:
            speakers.append(s["speaker"])
    return {"names": names, "speakers": [_speaker(s, names) for s in speakers]}


def export_docx(meeting: dict) -> Path:
    p = _prepare(meeting)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    h = doc.add_heading("ПРОТОКОЛ СОВЕЩАНИЯ", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(meeting["title"]).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Дата: {_fmt_date(meeting['meeting_date'])}")
    doc.add_paragraph(f"Участники: {', '.join(p['speakers']) or '—'}")
    if meeting.get("consent"):
        note = doc.add_paragraph(CONSENT_NOTE)
        note.runs[0].italic = True

    doc.add_heading("Краткое содержание", level=1)
    doc.add_paragraph(meeting.get("summary") or "—")

    if meeting.get("reports"):
        doc.add_heading("Саммари по ключевым пунктам", level=1)
        rt = doc.add_table(rows=1, cols=3)
        rt.style = "Table Grid"
        for cell, text in zip(rt.rows[0].cells, ["Направление / доклад", "Показатель", "Проблема"]):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for r in meeting["reports"]:
            row = rt.add_row().cells
            direction = r.get("direction") or "—"
            if r.get("speaker"):
                direction += f" ({_speaker(r['speaker'], p['names'])})"
            for cell, text in zip(row, [direction, r.get("metric") or "—", r.get("problem") or "—"]):
                cell.text = text

    if meeting.get("decisions"):
        doc.add_heading("Принятые решения", level=1)
        for d in meeting["decisions"]:
            doc.add_paragraph(d, style="List Number")

    doc.add_heading("Поручения", level=1)
    tasks = meeting.get("tasks", [])
    if tasks:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["№", "Поручение", "Ответственный", "Срок", "Приоритет"]):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for i, t in enumerate(tasks, 1):
            row = table.add_row().cells
            deadline = _fmt_date(t.get("deadline"))
            if t.get("deadline_text") and not t.get("deadline"):
                deadline = t["deadline_text"]
            for cell, text in zip(row, [str(i), t["task"], t["assignee_display"], deadline,
                                        t.get("priority") or ""]):
                cell.text = text
    else:
        doc.add_paragraph("Поручения не выявлены.")

    doc.add_heading("Стенограмма", level=1)
    for s in meeting.get("segments", []):
        para = doc.add_paragraph()
        r = para.add_run(f"[{_ts(s['start'])}] {_speaker(s['speaker'], p['names'])}: ")
        r.bold = True
        para.add_run(s["text"])

    foot = doc.add_paragraph("Протокол сформирован автоматически и подлежит проверке секретарём.")
    foot.runs[0].font.size = Pt(9)
    foot.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    path = EXPORT_DIR / f"protocol_{meeting['id']}.docx"
    doc.save(path)
    return path


_fonts_registered = False


def _register_fonts() -> None:
    global _fonts_registered
    if not _fonts_registered:
        # DejaVu поддерживает кириллицу, включая казахские буквы (ә, ғ, қ, ң, ө, ұ, ү, һ, і)
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONTS_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONTS_DIR / "DejaVuSans-Bold.ttf")))
        _fonts_registered = True


def export_pdf(meeting: dict) -> Path:
    _register_fonts()
    p = _prepare(meeting)
    body = ParagraphStyle("body", fontName="DejaVu", fontSize=10, leading=14)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11)
    muted = ParagraphStyle("muted", parent=body, fontSize=9, textColor=colors.HexColor("#555555"))
    title = ParagraphStyle("title", fontName="DejaVu-Bold", fontSize=16, leading=20, alignment=TA_CENTER)
    subtitle = ParagraphStyle("subtitle", parent=body, alignment=TA_CENTER, fontSize=11)
    h2 = ParagraphStyle("h2", fontName="DejaVu-Bold", fontSize=12, leading=16, spaceBefore=10, spaceAfter=4)
    e = lambda s: escape(str(s or ""))

    story = [Paragraph("ПРОТОКОЛ СОВЕЩАНИЯ", title), Paragraph(e(meeting["title"]), subtitle), Spacer(1, 6 * mm),
             Paragraph(f"<b>Дата:</b> {_fmt_date(meeting['meeting_date'])}", body),
             Paragraph(f"<b>Участники:</b> {e(', '.join(p['speakers']) or '—')}", body)]
    if meeting.get("consent"):
        story.append(Paragraph(f"<i>{CONSENT_NOTE}</i>", muted))

    story += [Paragraph("Краткое содержание", h2), Paragraph(e(meeting.get("summary") or "—"), body)]

    if meeting.get("reports"):
        story.append(Paragraph("Саммари по ключевым пунктам", h2))
        rrows = [[Paragraph(f"<b>{h}</b>", small) for h in ["Направление / доклад", "Показатель", "Проблема"]]]
        for r in meeting["reports"]:
            direction = r.get("direction") or "—"
            if r.get("speaker"):
                direction += f" ({_speaker(r['speaker'], p['names'])})"
            rrows.append([Paragraph(e(direction), small), Paragraph(e(r.get("metric") or "—"), small),
                          Paragraph(e(r.get("problem") or "—"), small)])
        rtable = Table(rrows, colWidths=[62 * mm, 45 * mm, 66 * mm], repeatRows=1)
        rtable.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(rtable)

    if meeting.get("decisions"):
        story.append(Paragraph("Принятые решения", h2))
        for i, d in enumerate(meeting["decisions"], 1):
            story.append(Paragraph(f"{i}. {e(d)}", body))

    story.append(Paragraph("Поручения", h2))
    tasks = meeting.get("tasks", [])
    if tasks:
        rows = [[Paragraph(f"<b>{h}</b>", small) for h in ["№", "Поручение", "Ответственный", "Срок", "Приоритет"]]]
        for i, t in enumerate(tasks, 1):
            deadline = _fmt_date(t.get("deadline"))
            if t.get("deadline_text") and not t.get("deadline"):
                deadline = t["deadline_text"]
            rows.append([Paragraph(str(i), small), Paragraph(e(t["task"]), small),
                         Paragraph(e(t["assignee_display"]), small), Paragraph(e(deadline), small),
                         Paragraph(e(t.get("priority")), small)])
        table = Table(rows, colWidths=[9 * mm, 80 * mm, 38 * mm, 24 * mm, 22 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    else:
        story.append(Paragraph("Поручения не выявлены.", body))

    story.append(Paragraph("Стенограмма", h2))
    for s in meeting.get("segments", []):
        story.append(Paragraph(
            f"<b>[{_ts(s['start'])}] {e(_speaker(s['speaker'], p['names']))}:</b> {e(s['text'])}", small))

    story += [Spacer(1, 6 * mm),
              Paragraph("Протокол сформирован автоматически и подлежит проверке секретарём.", muted)]

    path = EXPORT_DIR / f"protocol_{meeting['id']}.pdf"
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm,
                      title=f"Протокол: {meeting['title']}").build(story)
    return path
